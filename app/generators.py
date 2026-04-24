import os
import random
import string
import csv
from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

def random_text(size):
    return ''.join(
        random.choice(string.ascii_letters + string.digits + " \n")
        for _ in range(size)
    )

def generate_txt(path, size):
    with open(path, "w", encoding="utf-8") as f:
        written = 0
        while written < size:
            chunk = min(1024, size - written)
            text = random_text(chunk)
            f.write(text)
            written += len(text.encode("utf-8"))

def generate_csv(path, size):
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        while os.path.getsize(path) < size:
            writer.writerow([random_text(50), random_text(50)])

def generate_pdf(path, size):
    c = canvas.Canvas(path, pagesize=A4)
    y = 800
    while True:
        c.drawString(40, y, random_text(80))
        y -= 14
        if y < 40:
            c.showPage()
            y = 800
        c.save()
        if os.path.getsize(path) >= size:
            break

def generate_docx(path, size):
    doc = Document()
    while True:
        doc.add_paragraph(random_text(200))
        doc.save(path)
        if os.path.getsize(path) >= size:
            break

def generate_xlsx(path, size):
    wb = Workbook()
    ws = wb.active
    row = 1
    while True:
        ws.cell(row=row, column=1, value=random_text(100))
        row += 1
        wb.save(path)
        if os.path.getsize(path) >= size:
            break

def generate_bin(path, size):
    with open(path, "wb") as f:
        f.write(os.urandom(size))